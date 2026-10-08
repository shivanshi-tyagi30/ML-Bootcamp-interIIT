"""Word record export with python-docx (spec Section 10.14)."""

from __future__ import annotations

from pathlib import Path

from app.exports.txt import hms
from app.models.record import MeetingRecord


def _cite(ids: list[str]) -> str:
    return ""


def write_docx(record: MeetingRecord, path: Path) -> None:
    """Write the record to a .docx file with the same sections as the Markdown export."""
    from docx import Document

    r = record
    doc = Document()
    doc.add_heading(r.meta.title, level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Date: {r.meta.generated_at:%d %b %Y}    Duration: {hms(r.meta.duration_s)}\n")
    meta.add_run(f"Source file: {r.meta.source_file}")
    if r.meta.warnings:
        meta.add_run("\nWarnings: " + ", ".join(r.meta.warnings))

    doc.add_heading("Summary", level=1)
    for s in r.summary or []:
        doc.add_paragraph(f"{s.text}")
    if not r.summary:
        doc.add_paragraph("No summary was produced.")

    doc.add_heading("Minutes", level=1)
    for t in r.minutes:
        doc.add_heading(t.topic, level=2)
        for p in t.points:
            doc.add_paragraph(f"{p.text}", style="List Bullet")
    if not r.minutes:
        doc.add_paragraph("No minutes were produced.")

    doc.add_heading("Key decisions", level=1)
    for d in r.decisions:
        doc.add_paragraph(
            f"{d.id} {d.decision} (agreement: “{d.agreement_evidence}”)",
            style="List Bullet",
        )
    if not r.decisions:
        doc.add_paragraph("No decisions were reached.")

    doc.add_heading("Not agreed (open proposals)", level=1)
    for p in r.open_proposals:
        extra = " (no clear agreement)" if p.demoted_from_decision else ""
        doc.add_paragraph(f"{p.proposal}{extra}", style="List Bullet")
    if not r.open_proposals:
        doc.add_paragraph("None.")

    doc.add_heading("Action items", level=1)
    if r.action_items:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, h in zip(table.rows[0].cells, ["#", "Task", "Owner", "Deadline"]):
            cell.text = h
        for a in r.action_items:
            row = table.add_row().cells
            for cell, v in zip(row, [a.id, a.task, a.owner, a.deadline]):
                cell.text = v
    else:
        doc.add_paragraph("No action items were identified.")

    doc.save(str(path))
